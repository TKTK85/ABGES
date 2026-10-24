# -*- coding: utf-8 -*-
"""Compact, decision-oriented delivery files for genetic evaluations."""

import math
import os
from collections import defaultdict


def _get(record, *names):
    for name in names:
        if name in record:
            return record[name]
    lower = {str(key).lower(): value for key, value in record.items()}
    for name in names:
        if name.lower() in lower:
            return lower[name.lower()]
    return None


def _number(value):
    try:
        value = float(value)
        return value if math.isfinite(value) else None
    except (TypeError, ValueError):
        return None


def _format(value, digits=4):
    if value is None:
        return "—"
    if isinstance(value, float):
        return ("{:,.%df}" % digits).format(value)
    return str(value)


def _rows_for_trait(plan, context, trait):
    parsed = context.get("parsed_result") or {}
    if plan.get("mode") == "multi":
        records = (parsed.get("rand_by_trait") or {}).get(trait) or []
        columns = (parsed.get("rand_columns_by_trait") or {}).get(trait) or []
    else:
        records = parsed.get("rand") or []
        columns = parsed.get("rand_columns") or []
    return records, columns


def _records_to_rows(phenotype, records, trait, fixed_effects, genetic_variance=None,
                     genetic_reporting=None):
    by_id = {}
    for record in records:
        animal_id = _get(record, "ID", "id")
        if animal_id is None:
            continue
        key = str(animal_id).strip()
        by_id[key] = record

    rows = []
    def reliability(record):
        direct = _number(_get(record, "PA_R2", "R2", "reliability"))
        if direct is not None and 0.0 <= direct <= 1.0:
            return direct, "HIBLUP_PA_R2"
        pev = _number(_get(record, "PA_PEV", "PEV"))
        if pev is not None and genetic_variance and genetic_variance > 0:
            value = 1.0 - pev / genetic_variance
            if 0.0 <= value <= 1.0:
                return value, "1-PEV/VA"
        return None, "不可用"
    if hasattr(phenotype, "to_dict"):
        source = phenotype.to_dict(orient="records")
    else:
        source = list(phenotype or [])
    genetic_reporting = genetic_reporting or {}
    reference_column = genetic_reporting.get("reference_column")
    reference_value = genetic_reporting.get("reference_value")
    reference_mean = None
    reference_n = 0
    if reference_column and reference_value is not None:
        matched = []
        for item in source:
            if str(item.get(reference_column, "")).strip() != str(reference_value).strip():
                continue
            record = by_id.get(str(item.get("ID", item.get("id", ""))).strip(), {})
            value = _number(_get(record, "PA", "EBV", trait))
            if value is not None:
                matched.append(value)
        reference_n = len(matched)
        if reference_n >= int(genetic_reporting.get("minimum_n", 30) or 30):
            reference_mean = sum(matched) / reference_n
    seen = set()
    for item in source:
        animal_id = item.get("ID", item.get("id"))
        if animal_id is None:
            continue
        key = str(animal_id).strip()
        seen.add(key)
        record = by_id.get(key, {})
        phenotype_value = _number(item.get(trait))
        reliability_value, reliability_source = reliability(record)
        rows.append({
            "ID": key,
            "groups": {effect: item.get(effect) for effect in fixed_effects},
            "phenotype": phenotype_value,
            "phenotype_used": phenotype_value is not None,
            "EBV": _number(_get(record, "PA", "EBV", trait)),
            "EBV_reference": (_number(_get(record, "PA", "EBV", trait)) - reference_mean
                              if reference_mean is not None and _number(_get(record, "PA", "EBV", trait)) is not None else None),
            "EBV_SE": _number(_get(record, "PA_SE", "EBV_SE")),
            "PEV": _number(_get(record, "PA_PEV", "PEV")),
            "Reliability": reliability_value,
            "Reliability_source": reliability_source,
        })
    # Keep prediction-only pedigree animals in the result, even when they have
    # no row in the phenotype table supplied to the model.
    for key, record in by_id.items():
        if key in seen:
            continue
        reliability_value, reliability_source = reliability(record)
        rows.append({
            "ID": key,
            "groups": {effect: None for effect in fixed_effects},
            "phenotype": None,
            "phenotype_used": False,
            "EBV": _number(_get(record, "PA", "EBV", trait)),
            "EBV_reference": (_number(_get(record, "PA", "EBV", trait)) - reference_mean
                              if reference_mean is not None and _number(_get(record, "PA", "EBV", trait)) is not None else None),
            "EBV_SE": _number(_get(record, "PA_SE", "EBV_SE")),
            "PEV": _number(_get(record, "PA_PEV", "PEV")),
            "Reliability": reliability_value,
            "Reliability_source": reliability_source,
        })
    rows.sort(key=lambda row: (
        row["EBV"] is None,
        -(row["EBV"] or 0.0),
        row["ID"].casefold(),
    ))
    for rank, row in enumerate(rows, 1):
        row["Rank"] = rank
    return rows


def _qc_counts(context, traits):
    metrics = context.get("phenotype_metrics") or {}
    rows = metrics.get("rows") or 0
    change_count = 0
    for item in (context.get("phenotype_qc").change_log
                 if context.get("phenotype_qc") else []):
        try:
            change_count += int(item.get("affected") or 0)
        except (TypeError, ValueError):
            continue
    if not change_count and rows:
        # Do not call naturally missing values "excluded by QC"; this is the
        # explicit change-log count of records modified by a QC rule.
        change_count = 0
    return rows, change_count, metrics


def _group_rows(rows, fixed_effects, trait):
    grouped = defaultdict(list)
    for row in rows:
        for effect in fixed_effects:
            level = row["groups"].get(effect)
            if level is not None:
                grouped[(effect, str(level))].append(row)
    result = []
    for (effect, level), values in sorted(grouped.items()):
        phenotypes = [row["phenotype"] for row in values if row["phenotype"] is not None]
        ebvs = [row["EBV"] for row in values if row["EBV"] is not None]
        mean_pheno = sum(phenotypes) / len(phenotypes) if phenotypes else None
        sd_pheno = None
        if len(phenotypes) > 1:
            sd_pheno = math.sqrt(sum((x - mean_pheno) ** 2 for x in phenotypes)
                                 / (len(phenotypes) - 1))
        result.append([trait, effect, level, len(phenotypes), mean_pheno,
                       sd_pheno, len(ebvs),
                       sum(ebvs) / len(ebvs) if ebvs else None])
    return result


def _get_summary(validation, trait):
    summary = getattr(validation, "summary", None) or {}
    if isinstance(summary.get("heritability"), dict):
        return {
            "heritability": summary["heritability"].get(trait),
            "heritability_se": (summary.get("heritability_se") or {}).get(trait),
            "genetic_variance": (summary.get("genetic_variance") or {}).get(trait),
            "residual_variance": (summary.get("residual_variance") or {}).get(trait),
            "converged": summary.get("converged"),
            "n_iterations": summary.get("n_iterations"),
        }
    return summary


def _warnings(gate_warnings, stage_warnings):
    found = []
    for item in list(gate_warnings or []) + list(stage_warnings or []):
        message = item.get("message") or item.get("desc") or str(item)
        code = item.get("code") or item.get("check_id") or item.get("stage") or "告警"
        entry = "%s：%s" % (code, message)
        if entry not in found:
            found.append(entry)
    return found


def create_workbook(path, plan, context, validation, expected_status,
                    gate_warnings=None, stage_warnings=None, metadata=None):
    """Write an Excel workbook with key results, factor groups and full EBVs."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo

    phenotype_qc = context.get("phenotype_qc")
    phenotype = getattr(phenotype_qc, "frame", None)
    if phenotype is None:
        raise ValueError("缺少完成质控的表型数据，无法生成交付工作簿")
    traits = list(plan.get("traits") or [plan.get("trait")])
    fixed_effects = list(getattr(phenotype_qc, "effective_fixed_effects", []) or [])
    warnings = _warnings(gate_warnings, stage_warnings)
    rows_by_trait = {}
    for trait in traits:
        records, _ = _rows_for_trait(plan, context, trait)
        if not records:
            raise ValueError("性状 %s 缺少 EBV 明细，无法生成交付工作簿" % trait)
        summary = _get_summary(validation, trait)
        rows_by_trait[trait] = _records_to_rows(
            phenotype, records, trait, fixed_effects,
            genetic_variance=_number(summary.get("genetic_variance")),
            genetic_reporting=plan.get("genetic_reporting"))

    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "结果摘要"
    group_ws = wb.create_sheet("世代批次汇总")
    detail_ws = wb.create_sheet("育种值明细") if len(traits) == 1 else None
    navy = "17365D"
    blue = "D9EAF7"
    gold = "FFF2CC"

    def style_header(ws, row=1):
        for cell in ws[row]:
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 25

    def add_table(ws, name):
        if ws.max_row >= 2 and ws.max_column:
            ref = "A1:%s%d" % (ws.cell(1, ws.max_column).column_letter, ws.max_row)
            table = Table(displayName=name, ref=ref)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showFirstColumn=False,
                showLastColumn=False, showRowStripes=True, showColumnStripes=False)
            ws.add_table(table)

    summary_ws.append(["指标", "结果", "说明"])
    style_header(summary_ws)
    run_id = (metadata or {}).get("run_id") or context.get("__state__").run_id
    input_rows, qc_changed, metrics = _qc_counts(context, traits)
    for trait in traits:
        rowlist = rows_by_trait[trait]
        summary = _get_summary(validation, trait)
        observed = sum(1 for row in rowlist if row["phenotype_used"])
        summary_ws.append(["性状", trait, "原始量纲；EBV 按 HIBLUP 输出"])
        summary_ws.append(["交付状态", expected_status, "partial 表示存在告警，详见下方告警和运行日志"])
        summary_ws.append(["运行编号", run_id, "用于定位 workspace/_runs 中的完整审计记录"])
        summary_ws.append(["预测个体数", len(rowlist), "包括有表型和仅由系谱预测的个体"])
        summary_ws.append(["进入模型的有效表型数", observed, "PH/QC 后非缺失记录数"])
        summary_ws.append(["PH/QC 修改记录数", qc_changed, "按变更日志记录的排除/修改条数"])
        summary_ws.append(["固定效应", ", ".join(fixed_effects) or "无", "模型已拟合的固定效应"])
        summary_ws.append(["收敛", "是" if summary.get("converged") else "否/未知",
                           "迭代次数：%s" % (summary.get("n_iterations") or "—")])
        summary_ws.append(["遗传力 h²", _number(summary.get("heritability")),
                           "SE=%s" % _format(_number(summary.get("heritability_se")))])
        summary_ws.append(["加性遗传方差", _number(summary.get("genetic_variance")), "VA"])
        summary_ws.append(["残差方差", _number(summary.get("residual_variance")), "VE"])
        eval_result = context.get("result_evaluation") or {}
        if isinstance(eval_result, dict) and trait in eval_result:
            eval_result = eval_result[trait]
        reliability_eval = (eval_result.get("ebv_reliability") or {}) if isinstance(eval_result, dict) else {}
        if reliability_eval:
            summary_ws.append(["可靠度口径", reliability_eval.get("canonical_source") or "按个体来源列示",
                               "PA_R2 与 1-PEV/VA 最大绝对差=%s；异常值不纳入可靠度汇总" %
                               _format(_number(reliability_eval.get("r2_vs_pev_va_max_abs_difference")))])
        ebvs = [row["EBV"] for row in rowlist if row["EBV"] is not None]
        summary_ws.append(["EBV 均值", sum(ebvs) / len(ebvs) if ebvs else None,
                           "全量预测个体"])
        ref_cfg = plan.get("genetic_reporting") or {}
        ref_row = next((row for row in rowlist if row.get("EBV_reference") is not None), None)
        reference_mean = (ref_row["EBV"] - ref_row["EBV_reference"] if ref_row else None)
        summary_ws.append(["参考群 EBV 基准化", "已计算" if ref_row else "未计算",
                           ("%s=%s；参考群 EBV 均值=%s；个体列 EBV_reference=原始 EBV−参考群均值"
                            % (ref_cfg.get("reference_column"), ref_cfg.get("reference_value"),
                               _format(reference_mean)))
                           if ref_row else "任务未显式指定参考群，或参考群样本量不足；保留原始 EBV"])
        summary_ws.append(["EBV 范围", "%s 至 %s" % (
            _format(min(ebvs)) if ebvs else "—", _format(max(ebvs)) if ebvs else "—"),
            "同一性状内比较"])
        independent = context.get("independent_validation") or {}
        iv = independent.get(trait) or {}
        if iv:
            summary_ws.append(["独立验证状态", iv.get("status"),
                               "独立验证配对数=%s；Pearson 预测能力=%s；校准斜率=%s；RMSE=%s；预测能力不等同准确度" % (
                                   iv.get("n"), _format(_number(iv.get("pearson_predictive_ability"))),
                                   _format(_number(iv.get("calibration_slope"))),
                                   _format(_number(iv.get("rmse"))))])
        for warning in warnings:
            summary_ws.append(["告警", warning, "交付状态保留为 partial；请结合业务规则审阅"])
        summary_ws.append(["使用提示", "EBV 是模型估计值，不等同于直接观测表型。",
                           "育种决策前应结合候选群、亲缘、可靠度、健康和繁殖目标筛选。"])
    summary_ws.column_dimensions["A"].width = 25
    summary_ws.column_dimensions["B"].width = 34
    summary_ws.column_dimensions["C"].width = 68
    for row in summary_ws.iter_rows(min_row=2):
        if row[0].value == "告警":
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=gold)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    summary_ws.freeze_panes = "A2"

    group_ws.append(["性状", "分组因素", "水平", "有效表型数", "表型均值", "表型标准差",
                     "有 EBV 个体数", "EBV 均值"])
    for trait in traits:
        for row in _group_rows(rows_by_trait[trait], fixed_effects, trait):
            group_ws.append(row)
    style_header(group_ws)
    for col, width in enumerate([14, 18, 18, 17, 16, 16, 17, 16], 1):
        group_ws.column_dimensions[group_ws.cell(1, col).column_letter].width = width
    for row in group_ws.iter_rows(min_row=2, min_col=5, max_col=6):
        for cell in row:
            cell.number_format = "0.000"
    for row in group_ws.iter_rows(min_row=2, min_col=8, max_col=8):
        for cell in row:
            cell.number_format = "0.000"
    group_ws.freeze_panes = "A2"
    add_table(group_ws, "GroupSummary")

    if detail_ws is None:
        detail_ws = wb.create_sheet("育种值明细")
    detail_ws.append(["性状", "排名", "个体ID"] + fixed_effects + ["表型是否进入模型", "表型值",
                      "EBV（原始）", "EBV_reference", "EBV_SE", "PEV", "可靠度R²", "可靠度来源"])
    for trait in traits:
        for row in rows_by_trait[trait]:
            detail_ws.append([trait, row["Rank"], row["ID"]]
                             + [row["groups"].get(effect) for effect in fixed_effects]
                             + ["是" if row["phenotype_used"] else "否",
                                row["phenotype"], row["EBV"], row["EBV_reference"], row["EBV_SE"],
                                row["PEV"], row["Reliability"], row["Reliability_source"]])
    style_header(detail_ws)
    widths = [14, 10, 25] + [16] * len(fixed_effects) + [18, 16, 16, 16, 16, 16, 18, 22]
    for col, width in enumerate(widths, 1):
        detail_ws.column_dimensions[detail_ws.cell(1, col).column_letter].width = width
    detail_ws.freeze_panes = "D2"
    detail_ws.auto_filter.ref = detail_ws.dimensions
    for row in detail_ws.iter_rows(min_row=2, min_col=3 + len(fixed_effects),
                                   max_col=detail_ws.max_column):
        for cell in row:
            if isinstance(cell.value, (float, int)):
                cell.number_format = "0.0000"
    add_table(detail_ws, "EBVDetail")

    pedigree_qc = context.get("pedigree_qc")
    if pedigree_qc is not None:
        exclusions = list(getattr(pedigree_qc, "exclusions", []) or [])
        exclusion_ws = wb.create_sheet("系谱剔除明细")
        exclusion_ws.append(["原始行号", "ID", "父本ID", "母本ID", "世代", "批次", "性别",
                             "规则编号", "剔除类型", "剔除原因"])
        style_header(exclusion_ws)
        for item in exclusions:
            exclusion_ws.append([
                item.get("source_row"), item.get("ID"), item.get("Sire"), item.get("Dam"),
                item.get("Gen"), item.get("Batch"), item.get("Sex"), item.get("rule_id"),
                ("重复行去重" if item.get("record_disposition") == "duplicate_row_removed"
                 else "级联剔除" if item.get("cascade") else "直接剔除"), item.get("reason"),
            ])
        for col, width in enumerate([12, 20, 20, 20, 12, 16, 12, 14, 14, 54], 1):
            exclusion_ws.column_dimensions[exclusion_ws.cell(1, col).column_letter].width = width
        exclusion_ws.freeze_panes = "A2"
        exclusion_ws.auto_filter.ref = exclusion_ws.dimensions
        add_table(exclusion_ws, "PedigreeExclusions")
        pedigree_summary = getattr(pedigree_qc, "exclusion_summary", {}) or {}
        summary_ws.append(["系谱直接剔除数", pedigree_summary.get("direct_excluded_records", 0),
                           "见“系谱剔除明细”工作表"])
        summary_ws.append(["系谱级联剔除数", pedigree_summary.get("cascade_excluded_records", 0),
                           "删除错误个体及其后代，不删除其父母或兄弟姐妹"])
        summary_ws.append(["系谱剔除记录数", pedigree_summary.get("total_excluded_records", 0),
                           "包含重复行去重记录；直接个体剔除数单独列示"])
        summary_ws.append(["系谱完全重复行去重数", pedigree_summary.get("duplicate_rows_removed", 0),
                           "按 ID/Sire/Dam 三列完全相同判定"])

    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
    check = load_workbook(path, read_only=True, data_only=True)
    try:
        expected = {"结果摘要", "世代批次汇总", "育种值明细"}
        if not expected.issubset(set(check.sheetnames)):
            raise ValueError("生成的 Excel 缺少预期工作表")
        if check["育种值明细"].max_row != 1 + sum(map(len, rows_by_trait.values())):
            raise ValueError("生成的 Excel 个体明细行数与模型输出不一致")
    finally:
        check.close()
    return path


def render_markdown(path, plan, context, validation, expected_status,
                    gate_warnings=None, stage_warnings=None, metadata=None,
                    workbook_path=None):
    traits = list(plan.get("traits") or [plan.get("trait")])
    qc = context.get("phenotype_qc")
    fixed_effects = list(getattr(qc, "effective_fixed_effects", []) or [])
    metrics = context.get("phenotype_metrics") or {}
    rows, changed, _ = _qc_counts(context, traits)
    warnings = _warnings(gate_warnings, stage_warnings)
    run_id = (metadata or {}).get("run_id") or context.get("__state__").run_id
    lines = ["# %s 育种值交付摘要" % (plan.get("task_id") or run_id), "",
             "面向育种决策的核心结果与使用提示。全量个体明细、可靠度和分组筛选见 Excel；"
             "完整命令、日志、闸门和输入校验记录保存在本次运行目录。", "",
             "| 项目 | 结果 |", "|---|---|",
             "| 目标性状 | %s |" % ", ".join(traits),
             "| 模型 | %s |" % str(plan.get("model") or plan.get("model_id") or "遗传评估").upper(),
             "| 固定效应 | %s |" % (", ".join(fixed_effects) or "无"),
             "| 交付状态 | **%s** |" % expected_status,
             "| 运行编号 | %s |" % run_id,
             "| 输入记录数 | %s |" % (metrics.get("rows") if metrics.get("rows") is not None else rows),
             "| PH/QC 修改记录 | %s |" % changed,
             "| 有效表型记录 | %s |" % (metrics.get("valid_records") or "—"), ""]
    for trait in traits:
        summary = _get_summary(validation, trait)
        lines.extend(["## %s 核心估计" % trait, "",
                      "| 收敛 | 迭代次数 | h² | h² SE | VA | VE | EBV 个体数 |",
                      "|---|---:|---:|---:|---:|---:|---:|"])
        records, _ = _rows_for_trait(plan, context, trait)
        lines.append("| %s | %s | %s | %s | %s | %s | %s |" % (
            "是" if summary.get("converged") else "否/未知",
            summary.get("n_iterations") or "—",
            _format(_number(summary.get("heritability")), 6),
            _format(_number(summary.get("heritability_se")), 6),
            _format(_number(summary.get("genetic_variance")), 4),
            _format(_number(summary.get("residual_variance")), 4), len(records)))
        report_cfg = plan.get("genetic_reporting") or {}
        if report_cfg.get("reference_column") is not None:
            lines.append("显式参考群：**%s=%s**；个体 EBV_reference 仅在 Excel 记录参考群样本量达到阈值时提供。" %
                         (report_cfg.get("reference_column"), report_cfg.get("reference_value")))
        else:
            lines.append("未显式指定参考群；本次不生成基准化 EBV，原始 EBV 保持不变。")
        lines.append("")
        iv = (context.get("independent_validation") or {}).get(trait)
        if iv:
            lines.extend(["## 独立验证群预测表现", "",
                          "状态：**%s**；有效配对数：**%s**；Pearson 预测能力：**%s**；"
                          "Spearman：**%s**；校准斜率：**%s**；偏差：**%s**；RMSE：**%s**。"
                          % (iv.get("status"), iv.get("n"),
                             _format(_number(iv.get("pearson_predictive_ability"))),
                             _format(_number(iv.get("spearman_rank_correlation"))),
                             _format(_number(iv.get("calibration_slope"))),
                             _format(_number(iv.get("mean_bias_adjusted_phenotype_minus_ebv"))),
                             _format(_number(iv.get("rmse")))),
                          "该验证要求验证表型未参与拟合并已校正到模型目标尺度；预测能力指标不等于遗传评估准确度。", ""])
    pedigree_qc = context.get("pedigree_qc")
    pedigree_summary = getattr(pedigree_qc, "exclusion_summary", {}) or {}
    if pedigree_summary.get("total_excluded_records", 0):
        lines.extend(["## 系谱质控剔除", "",
                      "直接个体剔除：**%s** 条；级联剔除：**%s** 条；完全重复行去重：**%s** 条；记录处置合计：**%s** 条。"
                      % (pedigree_summary.get("direct_excluded_records", 0),
                         pedigree_summary.get("cascade_excluded_records", 0),
                         pedigree_summary.get("duplicate_rows_removed", 0),
                         pedigree_summary.get("total_excluded_records", 0)),
                      "其中与已通过表型 QC 个体重合的系谱剔除 ID：**%s**。"
                      % pedigree_summary.get("phenotype_individuals_affected", 0),
                      "逐个体 ID、父母 ID、规则、原因及世代/批次/性别见交付 Excel 的“系谱剔除明细”工作表。", ""])
        for group_name, levels in (pedigree_summary.get("by_group") or {}).items():
            lines.extend(["### 按%s剔除影响" % group_name, "",
                          "| 分组 | 输入记录 | 剔除记录 | 剔除比例 |",
                          "|---|---:|---:|---:|"])
            for level, values in sorted(levels.items()):
                lines.append("| %s | %s | %s | %.1f%% |" % (
                    level, values.get("input_records", 0), values.get("excluded_records", 0),
                    100.0 * values.get("excluded_rate", 0.0)))
            lines.append("")
    if warnings:
        lines.extend(["## 需要关注", ""])
        lines.extend("- %s" % warning for warning in warnings)
        lines.append("")
    else:
        lines.extend(["本次未发现需在交付摘要中提示的告警。", ""])
    if workbook_path:
        lines.extend(["## 交付文件", "", "- Excel：`%s`" % os.path.basename(workbook_path),
                      "- Excel 含结果摘要、世代/批次汇总、全部育种值明细。", ""])
    lines.extend(["## 解释边界", "",
                  "EBV 是模型在本次数据、模型和遗传基础上的估计值；有系谱但无有效表型的个体仍可能有 EBV。"
                  "排序只表示本模型内的估计值次序，不自动构成选种名单。请按候选群、可靠度、亲缘和育种目标复核。",
                  ""])
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    with open(path, "w", encoding="utf-8") as handle:
        handle.write("\n".join(lines))
    return path


def render_pedigree_qc_report(path, qc, species=None, generation_semantics=None):
    """Write a readable QC disposition report even when pedigree gates block evaluation."""
    summary = getattr(qc, "exclusion_summary", {}) or {}
    lines = ["# 系谱质量控制报告", "",
             "| 项目 | 数量/状态 |", "|---|---:|",
             "| 输入系谱记录 | %s |" % (qc.statistics.get("before", {}).get("Individuals", 0)),
             "| 直接剔除 | %s |" % summary.get("direct_excluded_records", 0),
             "| 级联剔除 | %s |" % summary.get("cascade_excluded_records", 0),
             "| 最终保留 | %s |" % summary.get("retained_records", 0),
             "| 完全重复行去重 | %s |" % summary.get("duplicate_rows_removed", 0),
             "| 最终系谱无环且引用闭合 | %s |" % ("是" if qc.final_graph_valid else "否"),
             "| 达到任务最小规模 | %s |" % ("是" if qc.minimum_size_passed else "否"),
             "| 物种 | %s |" % (species or "未指定"),
             "| 世代字段语义 | %s |" % (generation_semantics or "unspecified"), ""]
    lines.extend(["## 按规则汇总", "", "| 规则 | 剔除记录数 |", "|---|---:|"])
    for rule_id, count in sorted((summary.get("by_rule") or {}).items()):
        lines.append("| %s | %s |" % (rule_id, count))
    if summary.get("phenotype_individuals_affected") is not None:
        lines.extend(["", "表型 QC 后仍有记录、但因系谱 QC 被排除的个体数：**%s**。"
                      % summary.get("phenotype_individuals_affected", 0)])
    lines.extend(["", "## 按群体分层的剔除影响", ""])
    for group_name, levels in (summary.get("by_group") or {}).items():
        lines.extend(["### %s" % group_name, "",
                      "| 分组水平 | 输入记录 | 剔除记录 | 剔除比例 |",
                      "|---|---:|---:|---:|"])
        for level, values in sorted(levels.items()):
            lines.append("| %s | %d | %d | %.1f%% |" % (
                level, values.get("input_records", 0), values.get("excluded_records", 0),
                100.0 * values.get("excluded_rate", 0.0)))
        lines.append("")
    lines.extend(["## 逐条剔除记录", "",
                  "| 原始行 | ID | 父本 | 母本 | Gen | Batch | Sex | 规则 | 类型 | 原因 |",
                  "|---:|---|---|---|---|---|---|---|---|---|"])
    for item in qc.exclusions:
        values = [item.get("source_row"), item.get("ID"), item.get("Sire"), item.get("Dam"),
                  item.get("Gen"), item.get("Batch"), item.get("Sex"), item.get("rule_id"),
                  ("重复行去重" if item.get("record_disposition") == "duplicate_row_removed"
                   else "级联" if item.get("cascade") else "直接"), item.get("reason")]
        lines.append("| " + " | ".join("" if value is None else str(value).replace("|", "\\|")
                                       for value in values) + " |")
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    with open(path, "w", encoding="utf-8") as handle:
        handle.write("\n".join(lines))
    return path
