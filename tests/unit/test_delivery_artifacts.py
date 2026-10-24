from types import SimpleNamespace

import pandas as pd
from openpyxl import load_workbook

from runners.delivery_artifacts import create_workbook, render_markdown


def _context():
    qc = SimpleNamespace(
        frame=pd.DataFrame([
            {"ID": "A2", "Gen": 2, "Batch": 1, "EN": 10.0},
            {"ID": "A1", "Gen": 1, "Batch": 1, "EN": None},
        ]),
        effective_fixed_effects=["Gen", "Batch"],
        change_log=[{"rule_id": "PH006", "affected": 1}],
    )
    return {
        "phenotype_qc": qc,
        "phenotype_metrics": {"rows": 2, "valid_records": 1},
        "parsed_result": {"rand": [
            {"ID": "A2", "PA": 3.0, "PA_SE": 1.0, "PA_PEV": 1.0, "PA_R2": 0.8},
            {"ID": "A1", "PA": -1.0, "PA_SE": 1.5, "PA_PEV": 2.0, "PA_R2": 0.6},
            {"ID": "PRED_ONLY", "PA": 2.0, "PA_SE": 2.0, "PA_PEV": 4.0, "PA_R2": 0.4},
        ]},
        "__state__": SimpleNamespace(run_id="run_test"),
    }


def test_workbook_has_usable_summary_groups_and_all_ebvs(tmp_path):
    context = _context()
    validation = SimpleNamespace(summary={
        "trait": "EN", "heritability": 0.3, "heritability_se": 0.02,
        "genetic_variance": 4.0, "residual_variance": 10.0,
        "converged": True, "n_iterations": 8,
    })
    path = tmp_path / "育种值交付.xlsx"
    create_workbook(str(path), {
        "task_id": "task", "trait": "EN", "traits": ["EN"],
        "mode": "single", "model": "pblup",
    }, context, validation, "partial", [], [], {"run_id": "run_test"})

    workbook = load_workbook(path, read_only=False, data_only=True)
    assert workbook.sheetnames == ["结果摘要", "世代批次汇总", "育种值明细"]
    detail = workbook["育种值明细"]
    assert detail.max_row == 4
    headers = [cell.value for cell in detail[1]]
    phenotype_col = headers.index("表型是否进入模型") + 1
    id_col = headers.index("个体ID") + 1
    assert detail.cell(2, id_col).value == "A2"
    assert detail.cell(2, phenotype_col).value == "是"
    assert detail.cell(3, id_col).value == "PRED_ONLY"
    assert detail.cell(3, phenotype_col).value == "否"
    assert workbook["育种值明细"].auto_filter.ref is not None
    assert workbook["世代批次汇总"].max_row > 1
    workbook.close()


def test_markdown_is_compact_and_warns_about_partial_delivery(tmp_path):
    context = _context()
    validation = SimpleNamespace(summary={
        "heritability": 0.3, "heritability_se": 0.02,
        "genetic_variance": 4.0, "residual_variance": 10.0,
        "converged": True, "n_iterations": 8,
    })
    path = tmp_path / "summary.md"
    render_markdown(str(path), {
        "task_id": "task", "trait": "EN", "traits": ["EN"], "model": "pblup",
    }, context, validation, "partial", [], [{"code": "QC", "message": "需复核"}],
       {"run_id": "run_test"}, "delivery.xlsx")
    text = path.read_text(encoding="utf-8")
    assert "需复核" in text
    assert "PH/QC 修改记录 | 1" in text
    assert "delivery.xlsx" in text
    assert "command_line" not in text
